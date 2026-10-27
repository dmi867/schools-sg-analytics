#!/usr/bin/env python3
"""Merge Accent folder exports (Платежи / стройготовность) into data/per_object.json.

Expects layout:
  data/raw/accent_test/pay/<UIN>/accent_<UIN>_platezhi_*.xlsx
  data/raw/accent_test/sg/<UIN>/accent_<UIN>_stroygotovnost_*.xlsx

SG columns: Дата | Факт | План | Источник
Pay columns: Дата | … | Сумма | … | Тип платежа | …
Only UINs that have BOTH pay and sg folders are imported (matrix needs series).
"""
from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).parent.parent
OUT_PATH = ROOT / "data" / "per_object.json"
DEFAULT_ACCENT = ROOT / "data" / "raw" / "accent_test"


def parse_date(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if not v:
        return None
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(v).strip(), fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return None


def parse_num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace("\xa0", "").replace(" ", "").replace(",", "."))
    except ValueError:
        return None


def load_sg(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    series = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = parse_date(row[0] if row else None)
        if not d:
            continue
        fact = parse_num(row[1] if len(row) > 1 else None)
        plan = parse_num(row[2] if len(row) > 2 else None)
        if fact is None:
            continue
        series.append({"d": d, "plan": plan, "fact": fact})
    series.sort(key=lambda x: x["d"])
    return series


def load_pays(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    pays = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        d = parse_date(row[0])
        amt = parse_num(row[2] if len(row) > 2 else None)
        typ = str((row[4] if len(row) > 4 else "") or "").strip()
        if not d or amt is None:
            continue
        pays.append({"d": d, "amt": amt, "advance": typ == "Предоплата"})
    pays.sort(key=lambda x: x["d"])
    return pays


def latest_xlsx(folder: Path):
    files = sorted(folder.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    return files[0] if files else None


def main():
    accent = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_ACCENT
    pay_root = accent / "pay"
    sg_root = accent / "sg"
    if not pay_root.is_dir() or not sg_root.is_dir():
        sys.exit(f"Need {pay_root} and {sg_root}")

    existing = {}
    if OUT_PATH.exists():
        existing = json.loads(OUT_PATH.read_text(encoding="utf-8"))

    sg_uins = {p.name for p in sg_root.iterdir() if p.is_dir()}
    pay_uins = {p.name for p in pay_root.iterdir() if p.is_dir()}
    both = sorted(sg_uins & pay_uins)
    if not both:
        sys.exit("No UINs with both pay and sg")

    added = 0
    for uin in both:
        sg_f = latest_xlsx(sg_root / uin)
        pay_f = latest_xlsx(pay_root / uin)
        if not sg_f or not pay_f:
            print(f"SKIP {uin}: missing file")
            continue
        series = load_sg(sg_f)
        pays = load_pays(pay_f)
        if not series:
            print(f"SKIP {uin}: empty SG series")
            continue
        existing[uin] = {"series": series, "pays": pays}
        added += 1
        print(f"OK {uin}: series={len(series)} pays={len(pays)}")

    OUT_PATH.write_text(json.dumps(existing, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"Wrote {OUT_PATH} — {len(existing)} objects ({added} from Accent)")


if __name__ == "__main__":
    main()
