#!/usr/bin/env python3
"""Convert the per-object readiness/payment xlsx pairs (UIN.xlsx + UIN_платежи.xlsx) into a
single data/per_object.json. Kept so a future re-export of per-object files can be re-converted
the same way — build_html.py reads only the JSON, never these xlsx files directly."""
import json
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).parent.parent
SRC_DIR = ROOT if len(sys.argv) < 2 else Path(sys.argv[1])
OUT_PATH = ROOT / "data" / "per_object.json"


def parse_date(v):
    # Same formats build_html.py's own parse_date() accepts (datetime cell, or "%d.%m.%Y"/
    # "%Y-%m-%d" string) — kept in sync deliberately since this script's output must match
    # what build_html.py used to compute reading the xlsx files directly.
    dt = None
    if isinstance(v, datetime):
        dt = v
    elif v:
        for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
            try:
                dt = datetime.strptime(str(v).strip(), fmt)
                break
            except ValueError:
                pass
    return dt.strftime("%Y-%m-%d") if dt else None


def parse_amt(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace("\xa0", "").replace(" ", "").replace(",", "."))
    except ValueError:
        return None


def main():
    out = {}
    n_series = n_pays = 0
    for f in sorted(SRC_DIR.glob("*.xlsx")):
        if any(s in f.name for s in ("_платежи", "Акцент", "Simple", "КСГ", "ПИР", "Факт")):
            continue
        uin = f.stem
        wb = openpyxl.load_workbook(f, data_only=True)
        series = []
        for row in wb.active.iter_rows(min_row=2, values_only=True):
            d = parse_date(row[0])
            if not d:
                continue
            try:
                plan = float(row[1]) if row[1] is not None else None
                fact = float(row[2])
            except (TypeError, ValueError):
                continue
            series.append({"d": d, "plan": plan, "fact": fact})
        series.sort(key=lambda x: x["d"])
        if not series:
            continue
        n_series += 1

        pays = []
        pf = SRC_DIR / f"{uin}_платежи.xlsx"
        if pf.exists():
            wb2 = openpyxl.load_workbook(pf, data_only=True)
            hdr = [c.value for c in next(wb2.active.iter_rows(min_row=1, max_row=1))]
            if hdr and hdr[0] == "Тип платежа":
                for row in wb2.active.iter_rows(min_row=2, values_only=True):
                    d = parse_date(row[4])
                    amt = parse_amt(row[5])
                    if d and amt:
                        pays.append({"d": d, "amt": amt, "advance": row[0] == "Предоплата"})
            n_pays += 1
        pays.sort(key=lambda x: x["d"])

        out[uin] = {"series": series, "pays": pays}

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUT_PATH.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"OK: {OUT_PATH} — {len(out)} objects ({n_series} series files, {n_pays} payment files)")


if __name__ == "__main__":
    main()
